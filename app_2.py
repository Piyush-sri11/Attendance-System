from flask import Flask, render_template, Response, request, redirect, url_for
from flask_sqlalchemy import SQLAlchemy
from flask_socketio import SocketIO, emit
from face_recognition import face_encodings, load_image_file, compare_faces, face_locations, face_distance
import cv2
import numpy as np
import base64
from PIL import Image
from io import BytesIO
import os
from src.models import db,Person, Attendance, AttendanceStatus
from datetime import date
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from flask import send_file
import io

app = Flask(__name__, static_url_path='/static')
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///person.db'
db.init_app(app)
socketio = SocketIO(app)



@app.route('/')
def index():
    return render_template('index.html')

@app.route('/index1')
def index1():
    return render_template('index1.html')

@app.route('/index2')
def index2():
    return render_template('index2.html')

@app.route('/submit', methods=['POST'])
def submit():
    if request.method == 'POST':
        name = request.form['name']
        address = request.form['address']
        mobile = request.form['mobile']
        email = request.form['email']
        image = request.files['image']
        image_path = "static/images/" + image.filename
        image.save(image_path)  # Save image to a folder
        face_image = load_image_file(image_path)
        face_image = cv2.cvtColor(face_image, cv2.COLOR_BGR2RGB)
        face_location = face_locations(face_image)
        face_enc = face_encodings(face_image, face_location)[0]
        new_person = Person(name=name, address=address, mobile=mobile, email=email, image=image.filename, face_encoding=face_enc)
        db.session.add(new_person)
        db.session.commit()
        print("New person added to database:", new_person)  # Debugging statement
        return redirect(url_for('index1'))

@app.route('/stored_data')
def stored_data():
    persons = Person.query.all()
    return render_template('stored_data.html', persons=persons)

@app.route('/attendance')
def attendance():
    today = date.today()
    persons = Person.query.all()

    # Mark absent for persons without attendance today
    for person in persons:
        attendance_exists = Attendance.query.filter_by(
            person_id=person.id,
            date=today
        ).first()
        if not attendance_exists:
            absent_attendance = Attendance(
                person_id=person.id,
                date=today,
                status=AttendanceStatus.ABSENT
            )
            db.session.add(absent_attendance)
    db.session.commit()

    page = request.args.get('page', 1, type=int)
    per_page = 10  # Adjust as needed

    # Join Attendance and Person, order by name and date
    records = (
        db.session.query(Attendance, Person)
        .join(Person, Attendance.person_id == Person.id)
        .order_by(Person.name.asc(), Attendance.date.desc())
        .paginate(page=page, per_page=per_page)
    )

    # Prepare data for template
    attendance_records = [
        {
            'name': person.name,
            'date': attendance.date.strftime('%Y-%m-%d'),
            'status': attendance.status.value  # Use .value for string
        }
        for attendance, person in records.items
    ]

    return render_template(
        'attendance.html',
        attendance_records=attendance_records,
        pagination=records
    )

@app.route('/export_attendance')
def export_attendance():
    wb = Workbook()
    wb.remove(wb.active)  # Remove the default sheet

    persons = Person.query.order_by(Person.name).all()
    for person in persons:
        ws = wb.create_sheet(title=person.name)
        ws.append(['Date', 'Status'])
        records = Attendance.query.filter_by(person_id=person.id).order_by(Attendance.date.desc()).all()
        for record in records:
            ws.append([record.date.strftime('%Y-%m-%d'), record.status.value])

        # Optional: Set column widths
        ws.column_dimensions[get_column_letter(1)].width = 15
        ws.column_dimensions[get_column_letter(2)].width = 15

    # Save to a bytes buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name='attendance_records.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )

def process_frame(data):
    with app.app_context():
        known_encodings = []
        persons = Person.query.all()
        for person in persons:
            known_encodings.append(np.array(person.face_encoding))
        tolerance = 0.6
        base64_data = data['data']

        image_data = base64.b64decode(base64_data)
            
        image = np.array(Image.open(BytesIO(image_data)))
            
        small_frame = cv2.resize(image, (0, 0), fx=0.25, fy=0.25)
        small_frame = cv2.cvtColor(small_frame, cv2.COLOR_BGR2RGB)
        face_location = face_locations(small_frame)
        face_encoding = face_encodings(small_frame, face_location)

        for face in face_encoding:
            matches = compare_faces(known_encodings, face)
            face_distances = face_distance(known_encodings, face)
            best_match_index = np.argmin(face_distances)
            if matches[best_match_index]:
                matched_person = persons[best_match_index]
                print("Match found:", matched_person.name) 

                today = date.today()
                attendance = Attendance.query.filter_by(
                    person_id=matched_person.id,
                    date=today
                ).first()

                if attendance:
                    if attendance.status != AttendanceStatus.PRESENT:
                        attendance.status = AttendanceStatus.PRESENT
                        db.session.commit()
                        print(f"Attendance updated to PRESENT for {matched_person.name} on {today}")
                        socketio.emit('match_found', {
                            'name': matched_person.name,
                            'image': matched_person.image,
                            'message': 'Attendance updated to PRESENT!'
                        })
                    else:
                        print(f"Attendance already marked for {matched_person.name} today.")
                        socketio.emit('match_found', {
                            'name': matched_person.name,
                            'image': matched_person.image,
                            'message': 'Attendance already marked for today.'
                        })
                else:
                    new_attendance = Attendance(
                        person_id=matched_person.id,
                        date=today,
                        status=AttendanceStatus.PRESENT
                    )
                    db.session.add(new_attendance)
                    db.session.commit()
                    print(f"Attendance marked for {matched_person.name} on {today}")
                    socketio.emit('match_found', {
                        'name': matched_person.name,
                        'image': matched_person.image,
                        'message': 'Attendance marked successfully!'
                    })
            else:
                socketio.emit('no_match', {'message': "Person doesn't exist"})

@socketio.on('frame')
def handle_frame(data):
    process_frame(data)

if __name__ == '__main__':
    with app.app_context():
        db.create_all()
        print("Database tables created successfully")
    socketio.run(app)